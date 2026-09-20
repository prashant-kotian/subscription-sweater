#!/usr/bin/env python3
"""Offline self-test for the LLM Batch Tool.

Creates sample input files (.txt/.xlsx/.docx/.csv), checks the tool-policy
engine (web search / MCP / skills, directives, deny lists), checks terminal
mode (official Gemini CLI / Qwen Code clients — command building, MCP settings
merge, and a full run against a stub CLI), then runs a MOCK batch into each
output format and verifies the round-trip. No browser, no LLM, no network,
no real CLI installs needed:

    python selftest.py
"""
from __future__ import annotations

import os
import re
import sys
import tempfile

# Same Windows cp1252-vs-UTF8 fix as main.py -- this file's own final banner
# uses a checkmark that otherwise crashes the selftest with a false failure
# after every real check has already passed.
for _stream in (sys.stdout, sys.stderr):
    if hasattr(_stream, "reconfigure"):
        _stream.reconfigure(encoding="utf-8", errors="replace")

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from llm_batch import readers, writers  # noqa: E402
from llm_batch.engine import Engine, RunConfig  # noqa: E402
from llm_batch.policy import ToolPolicy, decide  # noqa: E402

SAMPLE_PROMPTS = [
    "What is the capital of France?",
    "Explain TCP in one short paragraph.\nKeep it under 50 words.",
    "List 3 fruits, one per line.",
    'Write a haiku about robots.',
]

SERVERS = {
    "github": {"command": "npx", "args": ["-y", "@modelcontextprotocol/server-github"]},
    "files": {"command": "npx", "args": ["-y", "@modelcontextprotocol/server-filesystem", "/tmp"]},
}


# --------------------------------------------------------------------------- #
def test_policy() -> None:
    quiet = lambda *a, **k: None  # noqa: E731

    # auto: web search only when the prompt looks real-time / lookup-y
    d = decide(ToolPolicy(web_search="auto"), "Search the web for today's gold price", quiet)
    assert d.web_search is True, "auto should detect web intent"
    d = decide(ToolPolicy(web_search="auto"), "What is the capital of France?", quiet)
    assert d.web_search is False, "auto should not force web search"

    # never wins over intent; always wins over nothing
    d = decide(ToolPolicy(web_search="never"), "search the web for X", quiet)
    assert d.web_search is False and d.force_off_web
    assert any("do NOT use web search" in i for i in d.instructions)
    d = decide(ToolPolicy(web_search="always"), "Write a haiku", quiet)
    assert d.web_search is True and any("use web search" in i for i in d.instructions)

    # per-prompt directives override the run policy (and are stripped)
    d = decide(ToolPolicy(web_search="always"),
               "@@websearch: off\nWhat's the weather in Mumbai?", quiet)
    assert d.web_search is False and d.force_off_web
    assert "@@" not in d.clean_prompt and "Mumbai" in d.clean_prompt

    # deny list is a hard constraint (beats directives + policy)
    d = decide(ToolPolicy(web_search="always", denied="websearch"),
               "@@websearch: on\nlatest news", quiet)
    assert d.web_search is False and d.force_off_web

    # MCP: always + allow list + per-server deny
    d = decide(ToolPolicy(mcp="always", mcp_servers=SERVERS), "do it", quiet)
    assert d.mcp and d.mcp_names == ["github", "files"]
    d = decide(ToolPolicy(mcp="always", mcp_servers=SERVERS, allowed_mcp="github"), "do it", quiet)
    assert d.mcp_names == ["github"]
    d = decide(ToolPolicy(mcp="always", mcp_servers=SERVERS, denied="mcp:github"), "do it", quiet)
    assert d.mcp and d.mcp_names == ["files"]
    d = decide(ToolPolicy(mcp="always", mcp_servers=SERVERS, denied="mcp"), "do it", quiet)
    assert d.mcp is False and d.force_off_mcp

    # MCP auto: on when the prompt mentions a configured server name or "mcp"
    d = decide(ToolPolicy(mcp="auto", mcp_servers=SERVERS), "push a commit via the github server", quiet)
    assert d.mcp is True and d.mcp_names == ["github", "files"]
    d = decide(ToolPolicy(mcp="auto", mcp_servers=SERVERS), "write a poem", quiet)
    assert d.mcp is False
    # directive selects a subset
    d = decide(ToolPolicy(mcp="always", mcp_servers=SERVERS),
               "@@mcp: use: files\ntask", quiet)
    assert d.mcp and d.mcp_names == ["files"]

    # enabled but nothing available -> off
    d = decide(ToolPolicy(mcp="always"), "use tools", quiet)
    assert d.mcp is False

    # skills: run-level + directive + deny
    d = decide(ToolPolicy(skills="pdf"), "summarize", quiet)
    assert d.skills == ["pdf"] and any("skill" in i.lower() for i in d.instructions)
    d = decide(ToolPolicy(skills="pdf"), "@@skill: excel\nsummarize", quiet)
    assert d.skills == ["pdf", "excel"]
    d = decide(ToolPolicy(skills="pdf", denied="skill:pdf"), "summarize", quiet)
    assert d.skills == []

    # no instructions when injection is off
    d = decide(ToolPolicy(web_search="always", inject_instructions=False), "x", quiet)
    assert d.web_search is True and d.instructions == []

    # model: run-level, per-prompt override, and empty default
    d = decide(ToolPolicy(model="GPT-4o"), "hi", quiet)
    assert d.target_model == "GPT-4o" and "model=GPT-4o" in d.summary()
    d = decide(ToolPolicy(model="GPT-4o"), "@@model: GPT-5\nhi", quiet)
    assert d.target_model == "GPT-5"
    d = decide(ToolPolicy(model="Qwen3-Max"), "@@model: qwen-long\nhi", quiet)
    assert d.target_model == "qwen-long"
    d = decide(ToolPolicy(), "hi", quiet)
    assert d.target_model == ""

    print("  policy  : web-search / MCP / skills / model decisions  OK")


# --------------------------------------------------------------------------- #
def test_marked(tmp: str) -> None:
    """Marked-block extraction: benchmark files wrap the real prompt between
    >>> PASTE BELOW >>> / <<< PASTE ABOVE <<< markers; only the marked part
    may go to the LLM, and the ID line above the marker is kept as a label."""
    mk = os.path.join(tmp, "marked.txt")
    with open(mk, "w", encoding="utf-8") as f:
        f.write(
            "GZ-01\n"
            "[reference only, do not paste -- source: secret gold data]\n\n"
            ">>> PASTE BELOW >>>\n"
            "Prompt one text.\n"
            "<<< PASTE ABOVE <<<\n\n"
            "==== BATCH HEADER — also must not be sent ====\n\n"
            "GZ-02\n"
            "[reference only, do not paste -- more secrets]\n\n"
            ">>> PASTE BELOW >>>\n"
            "Prompt two\n"
            "second line.\n"
            "<<< PASTE ABOVE <<<\n"
        )
    ps, labels = readers.read_prompts_labeled(mk)
    assert ps == ["Prompt one text.", "Prompt two\nsecond line."], ps
    assert labels == ["GZ-01", "GZ-02"], labels
    assert all("reference only" not in p for p in ps)
    assert all("BATCH HEADER" not in p for p in ps)

    # …and the IDs survive into a labeled txt output
    out = os.path.join(tmp, "marked_answers.txt")
    Engine(RunConfig(mode="mock", input_path=mk, output_path=out),
           log=lambda m: None, progress=lambda d, t: None).run()
    with open(out, encoding="utf-8") as f:
        content = f.read()
    assert "PROMPT 1 [GZ-01]" in content and "PROMPT 2 [GZ-02]" in content, content[:400]
    assert "secret gold data" not in content
    w = writers.make_writer(out, log=lambda m: None)
    assert w.count() == 2
    w.close()
    print("  marked  : marker extraction + labeled output  OK")


# --------------------------------------------------------------------------- #
def test_images(tmp: str) -> None:
    """Auto-attach by prompt label + explicit @@image: directive."""
    import os as _os

    d = _os.path.join(tmp, "imgs")
    _os.makedirs(d, exist_ok=True)
    for name in ("GZ-01_chart.png", "GZ-03_photo.jpg"):
        with open(_os.path.join(d, name), "wb") as f:
            f.write(b"\x89PNG fake")

    mk = _os.path.join(d, "prompts.txt")
    with open(mk, "w", encoding="utf-8") as f:
        f.write(
            "GZ-01\n>>> PASTE BELOW >>>\nChart question one.\n<<< PASTE ABOVE <<<\n\n"
            "GZ-02\n>>> PASTE BELOW >>>\nPlain text question.\n<<< PASTE ABOVE <<<\n\n"
            "GZ-03\n>>> PASTE BELOW >>>\n@@image: GZ-03_photo.jpg\nQuestion with directive.\n"
            "<<< PASTE ABOVE <<<\n"
        )
    logs: list = []
    out = _os.path.join(tmp, "img_answers.txt")
    Engine(RunConfig(mode="mock", input_path=mk, output_path=out, images_dir=d),
           log=logs.append, progress=lambda a, b: None).run()
    joined = "\n".join(logs)
    assert "(mock: image attached: GZ-01_chart.png)" in joined, joined
    assert "(mock: image attached: GZ-03_photo.jpg)" in joined, joined  # directive
    assert joined.count("image attached") == 2, joined  # GZ-02 gets none

    # explicit directive pointing at a missing file is announced, not fatal
    mk2 = _os.path.join(d, "prompts2.txt")
    with open(mk2, "w", encoding="utf-8") as f:
        f.write("GZ-09\n>>> PASTE BELOW >>>\n@@image: no_such_file.png\nQ\n<<< PASTE ABOVE <<<\n")
    logs2: list = []
    Engine(RunConfig(mode="mock", input_path=mk2, output_path=_os.path.join(tmp, "a2.txt"),
                     images_dir=d), log=logs2.append, progress=lambda a, b: None).run()
    assert "image not found: no_such_file.png" in "\n".join(logs2)
    print("  images  : auto-attach by label + @@image directive  OK")


# --------------------------------------------------------------------------- #
def test_writer_retry(tmp: str) -> None:
    """Re-running a prompt must REPLACE its entry (txt + xlsx), never stack a
    duplicate block on top of the old attempts."""
    out = os.path.join(tmp, "retry.txt")
    w = writers.make_writer(out, log=lambda m: None)
    w.add(1, "P-ONE", "first attempt (FAILED)")
    w.add(2, "P-TWO", "ok two")
    w.add(1, "P-ONE", "second attempt (final)")   # retry of prompt 1
    n = w.count()
    w.close()
    assert n == 2, f"retry: expected 2 blocks, got {n}"
    content = open(out, encoding="utf-8").read()
    assert content.count("PROMPT 1") == 1, content
    assert "second attempt (final)" in content
    assert "first attempt (FAILED)" not in content
    assert "ok two" in content                     # untouched block preserved
    # order preserved: prompt 1 still before prompt 2
    assert content.index("PROMPT 1") < content.index("PROMPT 2")

    # a file that ALREADY contains duplicate blocks (from the old append-only
    # behaviour) collapses to one entry for that index on the next write
    out2 = os.path.join(tmp, "retry_dup.txt")
    w2 = writers.TxtWriter(out2, log=lambda m: None)
    for r in ("attempt A", "attempt B", "attempt C"):
        w2._append(1, "P", r)                      # legacy behaviour -> 3 copies
    w2.add(1, "P", "final")
    n2 = w2.count()
    w2.close()
    assert n2 == 1, f"dedupe: expected 1 block, got {n2}"
    content2 = open(out2, encoding="utf-8").read()
    assert "final" in content2 and "attempt A" not in content2
    assert content2.count("PROMPT 1") == 1

    # image-chip artifact lines ("PNG" on its own line, as Gemini's web UI
    # copies them) are stripped; a sentence merely mentioning PNG survives
    outc = os.path.join(tmp, "chips.txt")
    wc = writers.make_writer(outc, log=lambda m: None)
    wc.add(1, "P", 'Based on the file "chart.png":\nPNG\n(a) 62 mN/m.\nPNG\n'
                   "Note: this is a PNG image, not a JPG.\n")
    wc.close()
    contentc = open(outc, encoding="utf-8").read()
    assert not re.search(r"(?m)^\s*PNG\s*$", contentc), contentc
    assert "this is a PNG image, not a JPG" in contentc
    assert "(a) 62 mN/m." in contentc

    # xlsx: re-running the same prompt ID replaces the row, no duplicate
    outx = os.path.join(tmp, "retry.xlsx")
    wx = writers.make_writer(outx, log=lambda m: None)
    wx.add(1, "P-ONE", "one v1", label="GZ-01")
    wx.add(2, "P-TWO", "two", label="GZ-02")
    wx.add(1, "P-ONE", "one v2", label="GZ-01")   # retry of GZ-01
    nx = wx.count()
    wx.close()
    assert nx == 2, f"xlsx retry: expected 2 rows, got {nx}"
    import openpyxl
    ws = openpyxl.load_workbook(outx).active
    vals = [(str(r[0]).strip(), str(r[2]).strip()) for r in ws.iter_rows(min_row=2, values_only=True)]
    assert ("GZ-01", "one v2") in vals, vals
    assert ("GZ-01", "one v1") not in vals, vals
    assert ("GZ-02", "two") in vals, vals
    print("  writers   : retry replaces entry (txt + xlsx), no duplicates  OK")


# --------------------------------------------------------------------------- #
def test_terminal(tmp: str, txt: str) -> None:
    import json as json_mod

    out_term = os.path.join(tmp, "answers_terminal.txt")
    from llm_batch import mcp_config
    from llm_batch.cli_bot import CLI_SITES, TerminalCLIBot
    from llm_batch.errors import ToolError
    from llm_batch.policy import PromptDecision

    quiet = lambda *a, **k: None  # noqa: E731

    # -- build_command: exact argv per policy decision ----------------------- #
    # client="gemini" pinned explicitly: select_client() auto-detects agy
    # over legacy gemini when agy is installed (the real, expected state on
    # most machines now that gemini-cli is EOL for individuals) -- these
    # assertions are specifically about the legacy gemini spec, so they must
    # not silently test whatever the running machine happens to have on PATH.
    bot = TerminalCLIBot(site="gemini", client="gemini", log=quiet)
    d = PromptDecision(clean_prompt="hello")
    cmd = bot.build_command("hello", d)
    assert cmd[:3] == ["gemini", "-p", "hello"] and "--yolo" in cmd, cmd

    d = PromptDecision(clean_prompt="hello", target_model="gemini-2.5-pro")
    assert bot.build_command("hello", d)[3:5] == ["--model", "gemini-2.5-pro"]

    d = PromptDecision(clean_prompt="hello", mcp=True, mcp_names=["github", "files"])
    cmd = bot.build_command("hello", d)
    i = cmd.index("--allowed-mcp-server-names")
    assert cmd[i + 1] == "github" and cmd.count("--allowed-mcp-server-names") == 2, cmd

    # mcp=OFF prompt in a run that HAS servers → allow-list matches nothing
    bot_srv = TerminalCLIBot(site="gemini", client="gemini", has_mcp_servers=True, log=quiet)
    d = PromptDecision(clean_prompt="hello", mcp=False)
    cmd = bot_srv.build_command("hello", d)
    i = cmd.index("--allowed-mcp-server-names")
    assert cmd[i + 1] == "__subscription_sweater_no_mcp__", cmd
    # …but a run with no servers configured → no allow-list flag at all
    cmd = TerminalCLIBot(site="qwen", log=quiet).build_command("hello",
                                                               PromptDecision(clean_prompt="h"))
    assert "--allowed-mcp-server-names" not in cmd, cmd

    # resume flag is real and site-specific: gemini's --resume documents
    # "latest" as a special value; qwen's --resume takes a real session ID
    # with no such case, so qwen uses the separate --continue flag instead.
    bot_nc_g = TerminalCLIBot(site="gemini", client="gemini", new_chat=False, log=quiet)
    assert bot_nc_g.build_command("h", PromptDecision(clean_prompt="h"))[-2:] == ["--resume", "latest"]
    bot_nc_q = TerminalCLIBot(site="qwen", new_chat=False, log=quiet)
    assert bot_nc_q.build_command("h", PromptDecision(clean_prompt="h"))[-1:] == ["--continue"]

    for s in CLI_SITES:
        assert TerminalCLIBot(site=s, log=quiet).spec["bin"] in (
            "codex", "claude", "qwen", "gemini", "agy")
    print("  terminal  : command builder (model / mcp scoping / resume)  OK")

    # -- Antigravity CLI (agy): gemini-cli's successor (retired 2026-06-18) -- #
    bot_agy = TerminalCLIBot(site="gemini", client="agy", new_chat=False, log=quiet)
    d = PromptDecision(clean_prompt="hello", target_model="gemini-3.5-flash",
                       mcp=True, mcp_names=["github"])
    cmd = bot_agy.build_command("hello", d)
    assert cmd[:3] == ["agy", "-p", "hello"], cmd
    assert cmd[3:5] == ["--model", "gemini-3.5-flash"], cmd
    assert ["--output-format", "json"] in [cmd[i:i + 2] for i in range(len(cmd) - 1)], cmd
    assert "--dangerously-skip-permissions" in cmd and "--yolo" not in cmd, cmd
    # agy has no allow-list flag — scoping happens via the mcp_config.json toggle
    assert "--allowed-mcp-server-names" not in cmd, cmd
    assert cmd[-1:] == ["--continue"], cmd   # official: -c/--continue = most recent

    # per-prompt MCP scoping: enabled run-servers present, the rest absent,
    # and the user's own servers preserved in both directions
    home2 = os.path.join(tmp, "fakehome2")
    cfgdir = os.path.join(home2, ".gemini", "config")
    os.makedirs(cfgdir, exist_ok=True)
    cfgp = os.path.join(cfgdir, "mcp_config.json")
    with open(cfgp, "w", encoding="utf-8") as f:
        json_mod.dump({"mcpServers": {"mine": {"command": "echo"}}}, f)
    # Windows' expanduser("~") reads USERPROFILE, not HOME -- both must be
    # faked (same pattern as the gemini MCP-merge block below) or this
    # silently resolves against the REAL home directory on Windows.
    old_home2 = os.environ.get("HOME"), os.environ.get("USERPROFILE")
    os.environ["HOME"] = home2
    if old_home2[1]:
        os.environ["USERPROFILE"] = home2
    try:
        assert mcp_config.cli_settings_path("agy") == os.path.normpath(cfgp), \
            mcp_config.cli_settings_path("agy")
        mcp_config.set_cli_mcp_scope("agy", SERVERS, {"github", "files"}, log=quiet)
        data = json_mod.load(open(cfgp, encoding="utf-8"))
        assert set(data["mcpServers"]) == {"mine", "github", "files"}, data
        mcp_config.set_cli_mcp_scope("agy", SERVERS, set(), log=quiet)
        data = json_mod.load(open(cfgp, encoding="utf-8"))
        assert set(data["mcpServers"]) == {"mine"}, data  # run servers gone, user's kept
    finally:
        os.environ["HOME"] = old_home2[0]
        if old_home2[1]:
            os.environ["USERPROFILE"] = old_home2[1]
    print("  terminal  : agy (Antigravity) command + per-prompt MCP scope  OK")

    # -- Codex CLI (OpenAI): site chatgpt -> `codex exec` --------------------- #
    bot_cx = TerminalCLIBot(site="chatgpt", log=quiet)
    assert bot_cx.client == "codex" and bot_cx.spec["bin"] == "codex"
    d = PromptDecision(clean_prompt="hello", target_model="gpt-5.6",
                       mcp=True, mcp_names=["github"])
    cmd = bot_cx.build_command("hello", d)
    assert cmd[:4] == ["codex", "exec", "--model", "gpt-5.6"], cmd
    assert "--yolo" in cmd and "--skip-git-repo-check" in cmd, cmd
    assert cmd[-1] == "hello", cmd                      # positional prompt, last
    i = cmd.index("--output-last-message")
    assert os.path.basename(cmd[i + 1]).startswith("sub-sweater-codex-"), cmd
    # mcp=on with an allowed subset -> the OTHER run servers get -c disabled
    bot_cx2 = TerminalCLIBot(site="chatgpt", mcp_servers=SERVERS, log=quiet)
    cmd = bot_cx2.build_command("h", d)
    assert ["-c", "mcp_servers.files.enabled=false"] in \
        [cmd[i:i + 2] for i in range(len(cmd) - 1)], cmd
    assert "mcp_servers.github.enabled=false" not in " ".join(cmd), cmd
    # mcp=off -> every run server disabled for this launch
    cmd = bot_cx2.build_command("h", PromptDecision(clean_prompt="h", mcp=False))
    assert "mcp_servers.github.enabled=false" in " ".join(cmd) and \
           "mcp_servers.files.enabled=false" in " ".join(cmd), cmd
    # continue = `exec resume --last`
    bot_cx3 = TerminalCLIBot(site="chatgpt", new_chat=False, log=quiet)
    cmd = bot_cx3.build_command("h", PromptDecision(clean_prompt="h"))
    assert cmd[:4] == ["codex", "exec", "resume", "--last"], cmd
    # native image attach
    cmd = bot_cx2.build_command("h", d, image_path=r"C:\imgs\c.png")
    assert ["--image", r"C:\imgs\c.png"] in \
        [cmd[i:i + 2] for i in range(len(cmd) - 1)], cmd
    print("  terminal  : codex (OpenAI) command + mcp overrides + image  OK")

    # -- Codex config.toml merge: append-only, comments + backup preserved --- #
    home_cx = os.path.join(tmp, "fakehome_codex")
    os.makedirs(os.path.join(home_cx, ".codex"), exist_ok=True)
    cp = os.path.join(home_cx, ".codex", "config.toml")
    with open(cp, "w", encoding="utf-8") as f:
        f.write('# user comment that must survive\nmodel = "gpt-5.6"\n\n'
                '[mcp_servers.mine]\ncommand = "echo"\n')
    old_home_cx = os.environ.get("HOME"), os.environ.get("USERPROFILE")
    os.environ["HOME"] = home_cx
    if old_home_cx[1]:
        os.environ["USERPROFILE"] = home_cx
    try:
        assert mcp_config.cli_settings_path("codex") == os.path.normpath(cp)
        mcp_config.ensure_cli_mcp_servers("codex", SERVERS, log=quiet)
        text = open(cp, encoding="utf-8").read()
        assert text.startswith("# user comment that must survive"), text
        assert 'model = "gpt-5.6"' in text
        assert "[mcp_servers.github]" in text and "[mcp_servers.files]" in text
        try:
            import tomllib
            data = tomllib.loads(text)
            assert set(data["mcp_servers"]) == {"mine", "github", "files"}
            assert data["mcp_servers"]["github"]["args"][1].endswith("server-github")
        except ImportError:
            pass
        backups = [x for x in os.listdir(os.path.join(home_cx, ".codex"))
                   if x.startswith("config.toml.bak-")]
        assert backups, "original config.toml must be backed up before merge"
        mcp_config.ensure_cli_mcp_servers("codex", SERVERS, log=quiet)  # idempotent
        text2 = open(cp, encoding="utf-8").read()
        assert text2.count("[mcp_servers.github]") == 1, "no duplicate sections"
    finally:
        os.environ["HOME"] = old_home_cx[0]
        if old_home_cx[1]:
            os.environ["USERPROFILE"] = old_home_cx[1]
    print("  terminal  : codex config.toml merge (append-only + backup)  OK")

    # -- Claude Code: site claude -> `claude -p`, strict per-launch mcp file -- #
    home_cc = os.path.join(tmp, "fakehome_claude")
    os.makedirs(home_cc, exist_ok=True)
    # the user's own server in ~/.claude.json must survive into the strict file
    with open(os.path.join(home_cc, ".claude.json"), "w", encoding="utf-8") as f:
        json_mod.dump({"other": 1, "mcpServers": {"mine": {"command": "echo"}}}, f)
    old_home_cc = os.environ.get("HOME"), os.environ.get("USERPROFILE")
    os.environ["HOME"] = home_cc
    if old_home_cc[1]:
        os.environ["USERPROFILE"] = home_cc
    try:
        bot_cc = TerminalCLIBot(site="claude", mcp_servers=SERVERS,
                                new_chat=False, log=quiet)
        assert bot_cc.client == "claude"
        d = PromptDecision(clean_prompt="hello", target_model="opus",
                           mcp=True, mcp_names=["files"])
        cmd = bot_cc.build_command("hello", d, image_path=r"C:\imgs\c.png")
        assert cmd[:2] == ["claude", "-p"], cmd
        # claude has no image flag: the path goes into the prompt text
        assert "C:\\imgs\\c.png" in cmd[2] and "file tools" in cmd[2], cmd[2]
        assert "--output-format" in cmd and "json" in cmd, cmd
        assert "--dangerously-skip-permissions" in cmd, cmd
        assert cmd[-1:] == ["--continue"], cmd
        i = cmd.index("--mcp-config")
        mcp_file = cmd[i + 1]
        assert cmd[i + 2] == "--strict-mcp-config", cmd
        data = json_mod.load(open(mcp_file, encoding="utf-8"))
        assert set(data["mcpServers"]) == {"mine", "files"}, data
        # mcp=off -> only the user's own servers remain in the strict file
        cmd = bot_cc.build_command("h", PromptDecision(clean_prompt="h", mcp=False))
        data = json_mod.load(open(mcp_file, encoding="utf-8"))
        assert set(data["mcpServers"]) == {"mine"}, data
    finally:
        os.environ["HOME"] = old_home_cc[0]
        if old_home_cc[1]:
            os.environ["USERPROFILE"] = old_home_cc[1]
    print("  terminal  : claude (Claude Code) command + strict mcp file  OK")

    # -- MCP settings merge (fake HOME, never the real one) ------------------ #
    home = os.path.join(tmp, "fakehome")
    os.makedirs(os.path.join(home, ".gemini"), exist_ok=True)
    gpath = os.path.join(home, ".gemini", "settings.json")
    with open(gpath, "w", encoding="utf-8") as f:
        json_mod.dump({"theme": "dark",
                       "mcpServers": {"existing": {"command": "echo"}}}, f)
    old_home = os.environ.get("HOME"), os.environ.get("USERPROFILE")
    os.environ["HOME"] = home
    if old_home[1]:
        os.environ["USERPROFILE"] = home
    try:
        assert mcp_config.cli_settings_path("gemini") == gpath
        assert mcp_config.cli_settings_path("qwen").endswith(os.path.join(".qwen", "settings.json"))
        p = mcp_config.ensure_cli_mcp_servers("gemini", SERVERS, log=quiet)
        assert p == gpath
        data = json_mod.load(open(gpath, encoding="utf-8"))
        assert data["theme"] == "dark"                        # user keys untouched
        assert set(data["mcpServers"]) == {"existing", "github", "files"}
        backups = [x for x in os.listdir(os.path.join(home, ".gemini"))
                   if x.startswith("settings.json.bak-")]
        assert backups, "original must be backed up before merge"
        mcp_config.ensure_cli_mcp_servers("gemini", SERVERS, log=quiet)  # idempotent
    finally:
        os.environ["HOME"] = old_home[0]
        if old_home[1]:
            os.environ["USERPROFILE"] = old_home[1]
    print("  terminal  : MCP settings merge (backup + user keys kept)  OK")

    # -- missing client → actionable error, not a crash ----------------------- #
    empty = os.path.join(tmp, "empty"); os.makedirs(empty, exist_ok=True)
    old_path = os.environ["PATH"]
    os.environ["PATH"] = empty
    try:
        try:
            TerminalCLIBot(site="gemini", log=quiet).start()
            raise AssertionError("expected ToolError for missing client")
        except ToolError as e:
            # guides to the CURRENT client (agy) and mentions the legacy fallback
            assert "antigravity.google/cli/install" in str(e), str(e)
            assert "npm install -g @google/gemini-cli" in str(e), str(e)
    finally:
        os.environ["PATH"] = old_path
    print("  terminal  : missing-client error message  OK")

    # -- end-to-end with a stub 'gemini' on PATH (posix only) ----------------- #
    if not (sys.platform.startswith("linux") or sys.platform == "darwin"):
        print("  terminal  : e2e stub skipped (posix only)")
        return
    bin_dir = os.path.join(tmp, "bin")
    os.makedirs(bin_dir, exist_ok=True)
    stub = os.path.join(bin_dir, "gemini")
    with open(stub, "w", encoding="utf-8") as f:
        f.write('#!/bin/sh\nprintf "%s\\n" "$@" > "$STUB_LOG"\n'
                'printf "STUB ANSWER: %s\\n" "$2"\n')
    os.chmod(stub, 0o755)
    stub_log = os.path.join(tmp, "stub_args.txt")
    os.environ["PATH"] = bin_dir + os.pathsep + old_path
    os.environ["STUB_LOG"] = stub_log
    os.environ["HOME"] = home                      # MCP merge must not touch the real one
    try:
        Engine(
            RunConfig(mode="terminal", site="gemini", input_path=txt,
                      output_path=out_term, web_search="never", mcp="always",
                      mcp_servers=SERVERS, delay_between=0.0),
            log=lambda m: None, progress=lambda d, t: None,
        ).run()
        w = writers.make_writer(out_term, log=lambda m: None)
        n = w.count(); w.close()
        assert n == len(SAMPLE_PROMPTS), f"terminal e2e: expected {len(SAMPLE_PROMPTS)}, got {n}"
        calls = open(stub_log, encoding="utf-8").read()
        assert "--yolo" in calls and "--allowed-mcp-server-names" in calls, calls
        g = json_mod.load(open(gpath, encoding="utf-8"))
        assert "github" in g["mcpServers"]

        # -- e2e with stub 'codex' (site chatgpt) and 'claude' ---------------- #
        cxstub = os.path.join(bin_dir, "codex")
        with open(cxstub, "w", encoding="utf-8") as f:
            f.write('#!/bin/sh\nprintf "%s\\n" "$@" >> "$STUB_LOG"\n'
                    'o=""; prev=""\nfor a in "$@"; do\n'
                    '  if [ "$prev" = "--output-last-message" ]; then o="$a"; fi\n'
                    '  prev="$a"\ndone\n'
                    '[ -n "$o" ] && printf "CODEX-OK" > "$o"\n'
                    'printf "CODEX STDOUT\\n"\n')
        os.chmod(cxstub, 0o755)
        cstub = os.path.join(bin_dir, "claude")
        with open(cstub, "w", encoding="utf-8") as f:
            f.write('#!/bin/sh\nprintf "%s\\n" "$@" >> "$STUB_LOG"\n'
                    'printf \'{"type":"result","is_error":false,'
                    '"result":"CLAUDE-OK","session_id":"s1"}\\n\'\n')
        os.chmod(cstub, 0o755)

        out_cx = os.path.join(tmp, "answers_codex.txt")
        Engine(RunConfig(mode="terminal", site="chatgpt", input_path=txt,
                         output_path=out_cx, web_search="never", mcp="always",
                         mcp_servers=SERVERS, delay_between=0.0),
               log=lambda m: None, progress=lambda d, t: None).run()
        w = writers.make_writer(out_cx, log=lambda m: None)
        assert w.count() == len(SAMPLE_PROMPTS)
        w.close()
        txt_cx = open(out_cx, encoding="utf-8").read()
        assert txt_cx.count("CODEX-OK") >= len(SAMPLE_PROMPTS), txt_cx
        cx_log = open(stub_log, encoding="utf-8").read()
        assert "codex" in cx_log and "--output-last-message" in cx_log, cx_log
        cx_cfg = os.path.join(home, ".codex", "config.toml")
        assert os.path.exists(cx_cfg) and \
            "[mcp_servers.github]" in open(cx_cfg, encoding="utf-8").read()

        out_cl = os.path.join(tmp, "answers_claude.txt")
        Engine(RunConfig(mode="terminal", site="claude", input_path=txt,
                         output_path=out_cl, web_search="never", mcp="always",
                         mcp_servers=SERVERS, delay_between=0.0),
               log=lambda m: None, progress=lambda d, t: None).run()
        txt_cl = open(out_cl, encoding="utf-8").read()
        assert txt_cl.count("CLAUDE-OK") >= len(SAMPLE_PROMPTS), txt_cl
        cl_cfg = os.path.join(home, ".claude.json")
        g2 = json_mod.load(open(cl_cfg, encoding="utf-8"))
        assert "github" in g2["mcpServers"]
        cl_log = open(stub_log, encoding="utf-8").read()
        assert "--mcp-config" in cl_log and "--strict-mcp-config" in cl_log
        print("  terminal  : e2e stub codex + claude (answers + configs)  OK")
    finally:
        os.environ["PATH"] = old_path
        os.environ["HOME"] = old_home[0]
        if old_home[1]:
            os.environ["USERPROFILE"] = old_home[1]
        os.environ.pop("STUB_LOG", None)
    print(f"  terminal  : e2e run via stub CLI ({n} answers, flags verified)  OK")


def make_inputs(tmp: str):
    txt = os.path.join(tmp, "prompts.txt")
    with open(txt, "w", encoding="utf-8") as f:
        f.write("\n\n".join(SAMPLE_PROMPTS))

    xlsx = os.path.join(tmp, "prompts.xlsx")
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["#", "Prompt"])
    for i, p in enumerate(SAMPLE_PROMPTS, 1):
        ws.append([i, p])
    wb.save(xlsx)

    docx = os.path.join(tmp, "prompts.docx")
    import docx as docx_mod

    d = docx_mod.Document()
    for p in SAMPLE_PROMPTS:
        d.add_paragraph(p)
    d.save(docx)

    csv = os.path.join(tmp, "prompts.csv")
    import csv as csv_mod

    with open(csv, "w", newline="", encoding="utf-8") as f:
        w = csv_mod.writer(f)
        w.writerow(["prompt"])
        for p in SAMPLE_PROMPTS:
            w.writerow([p])

    return txt, xlsx, docx, csv


def main() -> int:
    tmp = tempfile.mkdtemp(prefix="llm_batch_selftest_")
    print(f"sample inputs written to: {tmp}\n")

    # 1 · tool policy --------------------------------------------------------
    test_policy()
    test_marked(tmp)
    test_images(tmp)
    test_writer_retry(tmp)

    # 2 · reading ----------------------------------------------------------
    txt, xlsx, docx, csv = make_inputs(tmp)
    for name, src in [("txt", txt), ("xlsx", xlsx), ("docx", docx), ("csv", csv)]:
        got = readers.read_prompts(src)
        assert got == SAMPLE_PROMPTS, f"{name}: expected {len(SAMPLE_PROMPTS)} prompts, got {len(got)}: {got!r}"
        print(f"  read  {name:5s}: {len(got)} prompts  OK")

    # 3a · terminal mode (official Gemini CLI / Qwen Code clients) ----------- #
    test_terminal(tmp, txt)

    # 3 · mock run into every output format (with tool policy active) --------
    for ext in ["txt", "xlsx", "docx"]:
        out = os.path.join(tmp, f"answers.{ext}")
        Engine(
            RunConfig(mode="mock", input_path=txt, output_path=out,
                      web_search="always", mcp="always", mcp_servers=SERVERS,
                      skills="pdf"),
            log=lambda m: None, progress=lambda d, t: None,
        ).run()
        w = writers.make_writer(out, log=lambda m: None)
        n = w.count()
        w.close()
        assert n == len(SAMPLE_PROMPTS), f"{ext}: expected {len(SAMPLE_PROMPTS)} answers, got {n}"
        print(f"  write {ext:5s}: {n} answers (policy applied)  OK")

    # 4 · resume: a second run must not duplicate ----------------------------
    out = os.path.join(tmp, "answers.txt")
    Engine(RunConfig(mode="mock", input_path=txt, output_path=out),
           log=lambda m: None, progress=lambda d, t: None).run()
    w = writers.make_writer(out, log=lambda m: None)
    n = w.count()
    w.close()
    assert n == len(SAMPLE_PROMPTS), f"resume: expected {len(SAMPLE_PROMPTS)} answers, got {n}"
    print(f"  resume        : no duplicates  OK")

    print("\nSELFTEST PASSED ✔  (you can now run:  python main.py  to open the GUI)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
